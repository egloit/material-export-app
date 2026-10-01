"""
Excel Calculation Service for Material Export Application
Uses the formulas library to compute costs directly from supplier Excel files.
"""

import os
import logging
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

from config import Config

logger = logging.getLogger(__name__)


class ExcelCalculationService:
    """
    Computes packaging costs directly from supplier Excel files using the formulas library.
    Each .xlsx in the Excel/ folder represents one supplier.
    """

    def __init__(self, config: Config, app_logger=None):
        self.config = config
        self.app_logger = app_logger
        self._model_cache: Dict[str, object] = {}
        self._sheet_prefix_cache: Dict[str, str] = {}

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def list_suppliers(self) -> List[str]:
        """Scan base and upload folders, return supplier names (filename without .xlsx).        Upload folder takes precedence over base folder for duplicate names.
        Base suppliers listed in the blocklist are excluded unless overridden by an upload."""
        blocklist = self._get_blocklist()
        upload_folder = Path(self.config.EXCEL_FOLDER)
        uploaded = {
            Path(f).stem.lower()
            for f in os.listdir(upload_folder)
            if f.lower().endswith('.xlsx') and not f.startswith('~$')
        } if upload_folder.is_dir() else set()

        seen: Dict[str, None] = {}
        base_folder = Path(self.config.EXCEL_BASE_FOLDER)
        if base_folder.is_dir():
            for fname in sorted(os.listdir(base_folder)):
                if fname.lower().endswith('.xlsx') and not fname.startswith('~$'):
                    stem = Path(fname).stem
                    if stem.lower() not in blocklist or stem.lower() in uploaded:
                        seen[stem] = None
        if upload_folder.is_dir():
            for fname in sorted(os.listdir(upload_folder)):
                if fname.lower().endswith('.xlsx') and not fname.startswith('~$'):
                    seen[Path(fname).stem] = None
        return sorted(seen.keys())

    def find_matching_row_group(
        self,
        supplier: str,
        ply: str,
        box: str,
        size: str,
        design: str,
    ) -> Optional[int]:
        """
        Find the best-matching row group in the supplier's Excel file.

        Groups are identified by Sl.No. in column A: rows sharing the same
        Sl.No. value belong to the same calculation group.  The first row of
        each group is checked against the supplied criteria:
            H (col 8)  = Ply Type
            I (col 9)  = Box Type
            J (col 10) = Size Type
            K (col 11) = Design Type

        Returns the start row of the best match, or None if nothing matched.
        """
        excel_path = self._excel_path(supplier)
        if not excel_path:
            return None

        wb = openpyxl.load_workbook(str(excel_path), data_only=True, read_only=True)
        ws = wb.active

        # Build groups by Sl.No. (column A) – record the first row per Sl.No.
        groups: Dict[str, int] = {}  # sl_no -> first row
        max_row = ws.max_row or 200
        for row in range(6, max_row + 1):
            sl_no = ws.cell(row=row, column=1).value
            if sl_no is None:
                continue
            sl_key = str(sl_no).strip()
            if sl_key and sl_key not in groups:
                groups[sl_key] = row

        # Score each group by matching Ply/Box/Size/Design on its first row.
        # Box Type (col I) must match — skip groups where it doesn't.
        # Design Type (col K) must match too, but only if it is a design used in
        # the sheet (SAP descriptions can end in arbitrary letters, e.g. "GSM").
        known_designs = {
            str(ws.cell(row=r, column=11).value or '').strip().upper()
            for r in groups.values()
        }
        require_design = bool(design) and design.upper() in known_designs

        best_row = None
        best_score = -1

        for sl_key, first_row in groups.items():
            cell_h = str(ws.cell(row=first_row, column=8).value or '').strip()
            cell_i = str(ws.cell(row=first_row, column=9).value or '').strip()
            cell_j = str(ws.cell(row=first_row, column=10).value or '').strip()
            cell_k = str(ws.cell(row=first_row, column=11).value or '').strip()

            # Box Type is required — it must equal one of the "&"-separated
            # terms, so "KRT" does not match "KRTB" and "KM" matches "KM & HSC"
            box_terms = [t.strip().lower() for t in cell_i.split('&')]
            if box and box.lower() not in box_terms:
                continue
            if require_design and design.upper() != cell_k.upper():
                continue

            score = 0
            if ply and ply.lower() in cell_h.lower():
                score += 1
            if box and box.lower() in box_terms:
                score += 1
            # Prefer the plain box type ("KM") over combinations ("KM & HSC")
            if box and box.lower() == cell_i.lower():
                score += 1
            if size and size in cell_j:
                score += 1
            if design and design.upper() == cell_k.upper():
                score += 1

            if score > best_score:
                best_score = score
                best_row = first_row

        wb.close()

        if self.app_logger and best_row:
            self.app_logger.log_info(
                f"[ExcelService] {supplier}: best match row {best_row} (score {best_score}/5)"
            )

        return best_row

    def list_row_groups(self, supplier: str) -> List[Dict]:
        """
        Return all Sl.No. groups of a supplier's Excel file with the criteria
        of their first row: sl_no, row, desc (E), ply (H), box (I), size (J), design (K).
        Groups without Box Type (not applicable for this supplier) are skipped.
        """
        excel_path = self._excel_path(supplier)
        if not excel_path:
            return []

        wb = openpyxl.load_workbook(str(excel_path), data_only=True, read_only=True)
        ws = wb.active

        def text(v):
            return str(v).strip() if v is not None else ''

        groups = []
        seen = set()
        for row_idx, row in enumerate(ws.iter_rows(min_row=6, max_col=11, values_only=True), start=6):
            sl_key = text(row[0])
            if not sl_key or sl_key in seen:
                continue
            seen.add(sl_key)
            if not text(row[8]):
                continue
            groups.append({
                'sl_no': sl_key,
                'row': row_idx,
                'desc': text(row[4]),
                'ply': text(row[7]),
                'box': text(row[8]),
                'size': text(row[9]),
                'design': text(row[10]),
            })

        wb.close()
        return groups

    def list_all_row_groups(self) -> List[Dict]:
        """
        Merge the Sl.No. groups of all suppliers.  Criteria and description are
        taken from the first supplier that has the Sl.No.; 'rows' maps each
        supplier offering it to its start row.
        """
        merged: Dict[str, Dict] = {}
        for supplier in self.list_suppliers():
            for g in self.list_row_groups(supplier):
                entry = merged.setdefault(g['sl_no'], {**g, 'rows': {}})
                entry['rows'][supplier] = g['row']

        def sort_key(sl_no: str):
            try:
                return (0, float(sl_no))
            except ValueError:
                return (1, sl_no)

        return [merged[k] for k in sorted(merged, key=sort_key)]

    def files_signature(self) -> Tuple:
        """(supplier, path, mtime) of all active supplier files – changes on upload/delete."""
        sig = []
        for supplier in self.list_suppliers():
            p = self._excel_path(supplier)
            if p:
                sig.append((supplier, str(p), p.stat().st_mtime))
        return tuple(sig)

    def compute_all_suppliers_for_group(
        self,
        rows: Dict[str, int],
        L: float,
        B: float,
        H: float,
    ) -> Dict[str, Optional[Dict]]:
        """
        Compute costs for an explicitly chosen Sl.No. group.

        rows maps supplier -> start row (see list_all_row_groups); suppliers
        without that group get None.
        """
        results = {}
        for supplier in self.list_suppliers():
            start_row = rows.get(supplier)
            if start_row is None:
                results[supplier] = None
                continue
            try:
                results[supplier] = self.compute_from_excel(supplier, start_row, L, B, H)
            except Exception as e:
                if self.app_logger:
                    self.app_logger.log_error("ExcelService", f"{supplier}: {e}")
                results[supplier] = None
        return results

    def _get_group_size(self, supplier: str, start_row: int) -> int:
        """Determine how many rows belong to the Sl.No. group starting at start_row."""
        excel_path = self._excel_path(supplier)
        if not excel_path:
            return 3  # fallback

        wb = openpyxl.load_workbook(str(excel_path), data_only=True, read_only=True)
        ws = wb.active
        sl_no = ws.cell(row=start_row, column=1).value
        if sl_no is None:
            wb.close()
            return 3

        count = 0
        r = start_row
        max_row = ws.max_row or 200
        while r <= max_row:
            if str(ws.cell(row=r, column=1).value or '').strip() == str(sl_no).strip():
                count += 1
                r += 1
            else:
                break

        wb.close()
        return count if count > 0 else 3

    def compute_from_excel(
        self,
        supplier: str,
        start_row: int,
        L: float,
        B: float,
        H: float,
    ) -> Optional[Dict]:
        """
        Compute AJ values and AK for a given supplier, start_row and dimensions.

        The start_row is the first row of a Sl.No. group.  Dimensions (L/B/H)
        are injected into that row, and the group size is determined dynamically
        from the Sl.No. column.

        Returns dict with keys: 'AJ' (list of floats), 'AK' (float), 'rows' (detail list)
        or None on error.
        """
        try:
            import formulas as _formulas  # noqa: F811
        except ImportError:
            if self.app_logger:
                self.app_logger.log_error("ExcelService", "formulas library not installed (pip install formulas)")
            return None

        model = self._load_model(supplier)
        if model is None:
            return None

        sheet_prefix = self._detect_sheet_prefix(model, supplier)
        group_size = self._get_group_size(supplier, start_row)

        # Set L/B/H into the first row of the group
        inputs = {
            f"{sheet_prefix}!L{start_row}": L,
            f"{sheet_prefix}!M{start_row}": B,
            f"{sheet_prefix}!N{start_row}": H,
        }

        try:
            solution = model.calculate(inputs=inputs)
        except Exception as e:
            if self.app_logger:
                self.app_logger.log_error("ExcelService", f"{supplier} calculate failed: {e}")
            return None

        # Extract AJ for each row in the group and AK from the first row
        AJ = []
        rows_detail = []
        for i in range(group_size):
            r = start_row + i
            aj_val = self._extract_value(solution, f"{sheet_prefix}!AJ{r}")
            AJ.append(float(aj_val) if aj_val is not None else 0.0)

            rows_detail.append({
                'row': r,
                'T': self._extract_value(solution, f"{sheet_prefix}!T{r}"),
                'U': self._extract_value(solution, f"{sheet_prefix}!U{r}"),
                'V': self._extract_value(solution, f"{sheet_prefix}!V{r}"),
                'W': self._extract_value(solution, f"{sheet_prefix}!W{r}"),
                'Z': self._extract_value(solution, f"{sheet_prefix}!Z{r}"),
                'AF': self._extract_value(solution, f"{sheet_prefix}!AF{r}"),
                'AJ': AJ[-1],
            })

        ak_val = self._extract_value(solution, f"{sheet_prefix}!AK{start_row}")
        AK = float(ak_val) if ak_val is not None else sum(AJ)

        return {
            'AJ': AJ,
            'AK': AK,
            'rows': rows_detail,
            'start_row': start_row,
        }

    def compute_all_suppliers(
        self,
        ply: str,
        box: str,
        size: str,
        design: str,
        L: float,
        B: float,
        H: float,
    ) -> Dict[str, Optional[Dict]]:
        """
        Iterate over ALL supplier Excel files, find the matching row group,
        compute costs and return results keyed by supplier name.

        Returns: {'SKE': {'AJ': [...], 'AK': ...}, 'Ujala': {'AJ': [...], 'AK': ...}, ...}
        Suppliers that fail or have no match get None as value.
        """
        results = {}
        for supplier in self.list_suppliers():
            try:
                start_row = self.find_matching_row_group(supplier, ply, box, size, design)
                if start_row is None:
                    if self.app_logger:
                        self.app_logger.log_warning(
                            "ExcelService",
                            f"{supplier}: no matching row group for ply={ply}, box={box}, size={size}, design={design}"
                        )
                    results[supplier] = None
                    continue

                result = self.compute_from_excel(supplier, start_row, L, B, H)
                results[supplier] = result

            except Exception as e:
                if self.app_logger:
                    self.app_logger.log_error("ExcelService", f"{supplier}: {e}")
                results[supplier] = None

        return results

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    def _excel_path(self, supplier: str) -> Optional[Path]:
        """Return the full path to a supplier's Excel file, or None.
        Upload folder (EXCEL_FOLDER) is checked first; falls back to EXCEL_BASE_FOLDER.
        Base files in the blocklist are skipped unless overridden by an upload."""
        blocklist = self._get_blocklist()
        folders = [
            (self.config.EXCEL_FOLDER, False),
            (self.config.EXCEL_BASE_FOLDER, True),
        ]
        for folder_str, is_base in folders:
            if is_base and supplier.lower() in blocklist:
                continue
            folder = Path(folder_str)
            p = folder / f"{supplier}.xlsx"
            if p.exists():
                return p
            if folder.is_dir():
                for f in folder.iterdir():
                    if f.suffix.lower() == '.xlsx' and f.stem.lower() == supplier.lower():
                        return f
        return None

    def _get_blocklist(self) -> set:
        """Return the set of blocked base supplier names (lowercase)."""
        bl_file = Path(self.config.EXCEL_FOLDER) / ".blocklist"
        if not bl_file.exists():
            return set()
        return {line.strip().lower() for line in bl_file.read_text().splitlines() if line.strip()}

    def block_base_supplier(self, supplier: str) -> None:
        """Add a base supplier to the blocklist."""
        bl_file = Path(self.config.EXCEL_FOLDER) / ".blocklist"
        blocked = self._get_blocklist()
        blocked.add(supplier.lower())
        bl_file.write_text("\n".join(sorted(blocked)))
        self._model_cache.pop(supplier.lower(), None)

    def unblock_base_supplier(self, supplier: str) -> None:
        """Remove a base supplier from the blocklist."""
        bl_file = Path(self.config.EXCEL_FOLDER) / ".blocklist"
        blocked = self._get_blocklist()
        blocked.discard(supplier.lower())
        bl_file.write_text("\n".join(sorted(blocked)))

    def _load_model(self, supplier: str):
        """Load and cache the formulas ExcelModel for a supplier."""
        import formulas as _formulas

        cache_key = supplier.lower()
        if self.config.EXCEL_CACHE_MODELS and cache_key in self._model_cache:
            return self._model_cache[cache_key]

        excel_path = self._excel_path(supplier)
        if not excel_path:
            if self.app_logger:
                self.app_logger.log_error("ExcelService", f"Excel file not found for supplier: {supplier}")
            return None

        try:
            if self.app_logger:
                self.app_logger.log_info(f"[ExcelService] Loading model for {supplier}...")
            model = _formulas.ExcelModel().loads(str(excel_path)).finish()
            if self.config.EXCEL_CACHE_MODELS:
                self._model_cache[cache_key] = model
            return model
        except Exception as e:
            if self.app_logger:
                self.app_logger.log_error("ExcelService", f"Failed to load model for {supplier}: {e}")
            return None

    def _detect_sheet_prefix(self, model, supplier: str) -> str:
        """
        Detect the sheet prefix format from the model's solution keys.
        formulas uses formats like "'[SKE.xlsx]TABELLE1'" or "'Tabelle1'".
        """
        cache_key = supplier.lower()
        if cache_key in self._sheet_prefix_cache:
            return self._sheet_prefix_cache[cache_key]

        # Calculate with defaults to get sample keys
        try:
            sample_solution = model.calculate()
            sample_keys = list(sample_solution.keys())
        except Exception:
            sample_keys = []

        # Collect all unique sheet prefixes from the solution keys
        all_prefixes = set()
        for key in sample_keys:
            if '!' in key:
                all_prefixes.add(key.split('!')[0])

        # 1) Try to find a prefix whose sheet name matches the supplier name
        #    e.g. supplier="Ujala" should match "'[Ujala.xlsx]UJALA_V3'"
        #    Also try to match the active sheet name from openpyxl
        active_sheet = None
        excel_path = self._excel_path(supplier)
        if excel_path:
            try:
                wb = openpyxl.load_workbook(str(excel_path), data_only=True, read_only=True)
                active_sheet = wb.active.title.upper()
                wb.close()
            except Exception:
                pass

        prefix = "'Tabelle1'"  # default fallback
        supplier_matches = []
        for candidate in sorted(all_prefixes):
            upper = candidate.upper()
            # Match by active sheet name (most reliable)
            if active_sheet and active_sheet in upper:
                supplier_matches.append(candidate)
            # Match by supplier name
            elif supplier.upper() in upper:
                supplier_matches.append(candidate)

        if supplier_matches:
            # Prefer the highest version (last when sorted, e.g. V3 > V2 > V1)
            prefix = supplier_matches[-1]
        else:
            # Fallback: prefer any prefix with '[' (file reference), else first found
            for candidate in sorted(all_prefixes):
                if '[' in candidate:
                    prefix = candidate
                    break
                if prefix == "'Tabelle1'":
                    prefix = candidate

        self._sheet_prefix_cache[cache_key] = prefix

        if self.app_logger:
            self.app_logger.log_info(f"[ExcelService] {supplier}: detected sheet prefix = {prefix}")

        return prefix

    def _extract_value(self, solution: dict, cell_ref: str):
        """
        Extract a scalar value from the formulas solution dict.
        Handles numpy arrays, Ranges objects, etc.
        """
        val = solution.get(cell_ref)
        if val is None:
            # Try uppercase variant
            val = solution.get(cell_ref.upper())
        if val is None:
            return None

        # Unpack numpy / Ranges wrappers
        if hasattr(val, 'value'):
            val = val.value
        if hasattr(val, 'tolist'):
            val = val.tolist()
            while isinstance(val, list) and len(val) == 1:
                val = val[0]
        if hasattr(val, 'item'):
            val = val.item()

        return val
